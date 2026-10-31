"""Tests hors ligne de la planche : regroupement, mise en page, fichier .pptx (photos fictives)."""
import io
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "collector"))

import analyze  # noqa: E402
import board  # noqa: E402
from PIL import Image  # noqa: E402


def jpeg(w=300, h=450, color=(120, 100, 90)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "JPEG")
    return buf.getvalue()


def pant(n, brand="A", shape="droit", rise="haute", style="décontracté", closure="", day=None, **kw):
    item = {"id": f"{brand}:{n}", "brand": brand, "name": f"Pantalon {n}", "type": "Pantalons",
            "image": f"https://cdn.example/{brand}{n}.jpg", "url": f"https://x.example/p/{n}",
            "price": "100", "currency": "CAD", "first_seen": day or "2099-01-01T00:00:00+00:00",
            "has_vision": True, "attrs_v": analyze.ATTRS_VERSION,
            "shape": shape, "rise": rise, "style": style, "closure": closure}
    item.update(kw)
    return item


class GroupTests(unittest.TestCase):
    def test_cells_cover_grid_without_overlap(self):
        boxes = board.cell_boxes()
        self.assertEqual(len(boxes), 11)
        for x, y, w, h in boxes:
            self.assertGreaterEqual(x, 0)
            self.assertLessEqual(x + w, board.SLIDE_W + 1e-9)
            self.assertLessEqual(y + h, board.SLIDE_H + 1e-9)
        big = boxes[0]
        self.assertAlmostEqual(big[3], board.GRID_H)

    def test_full_group_first_then_widened(self):
        items = [pant(i, brand="AB"[i % 2]) for i in range(11)]  # 11 identiques : une diapo complète
        items += [pant(20 + i, brand="C", shape="large", rise="basse", style="habillé") for i in range(4)]
        items += [pant(30 + i, brand="D", shape="large", rise="basse", style="tailleur") for i in range(3)]
        slides = board.build_groups(items)
        full = [s for s in slides if not s["partial"]]
        self.assertEqual(len(full), 1)
        self.assertEqual(len(full[0]["items"]), 11)
        self.assertIn("DROIT", full[0]["label"])
        self.assertIn("TAILLE HAUTE", full[0]["label"])
        # les 7 « large / basse » se regroupent en élargissant (style retiré)
        partial = [s for s in slides if s["partial"]]
        self.assertEqual(len(partial), 1)
        self.assertEqual(len(partial[0]["items"]), 7)
        self.assertNotIn("HABILLÉ", partial[0]["label"])

    def test_tiny_groups_are_dropped_and_big_split(self):
        few = [pant(1, shape="jogger"), pant(2, shape="jogger")]
        self.assertEqual(board.build_groups(few), [])
        many = [pant(i, brand="AB"[i % 2]) for i in range(24)]
        slides = board.build_groups(many)
        self.assertEqual([len(s["items"]) for s in slides], [11, 11])  # le reste (2) est trop petit
        self.assertTrue(all("(" in s["label"] for s in slides))

    def test_no_item_used_twice(self):
        items = [pant(i, shape=("droit", "large")[i % 2], rise=("haute", "basse")[i % 3 == 0]) for i in range(40)]
        slides = board.build_groups(items)
        ids = [i["id"] for s in slides for i in s["items"]]
        self.assertEqual(len(ids), len(set(ids)))

    def test_interleave_mixes_brands(self):
        items = [pant(i, brand="A") for i in range(5)] + [pant(10 + i, brand="B") for i in range(5)]
        first = [i["brand"] for i in board.interleave(items)[:4]]
        self.assertEqual(sorted(first), ["A", "A", "B", "B"])

    def test_candidates_filter(self):
        ok = pant(1)
        old = pant(2, day="2000-01-01T00:00:00+00:00")
        no_photo_read = pant(3, has_vision=False)
        stale = pant(4, attrs_v=1)
        bad_url = pant(5, image="file:///etc/passwd")
        no_shape = pant(6, shape="")
        acc = pant(7, type="Accessoires")
        got = board.candidates([ok, old, no_photo_read, stale, bad_url, no_shape, acc], days=30)
        self.assertEqual([i["id"] for i in got], [ok["id"]])
        self.assertEqual(len(board.candidates([ok, old], days=0)), 2)


class DeckTests(unittest.TestCase):
    def build(self, n, bad=()):
        items = [pant(i, brand="AB"[i % 2], name=f"<b>Pantalon {i}</b>") for i in range(n)]
        out = Path(tempfile.mkdtemp()) / "b.pptx"

        def getter(url):
            if any(b in url for b in bad):
                raise OSError("hors ligne")
            return jpeg()

        result = board.generate(items, "Big ideas - Q2 2027", "Marque", 30, {"rates": {}}, out, getter=getter, sleep=lambda s: None)
        return result, out

    def test_full_deck(self):
        from pptx import Presentation

        result, out = self.build(11)
        self.assertEqual((result["slides"], result["partial"]), (1, 0))
        slide = Presentation(str(out)).slides[0]
        pics = [s for s in slide.shapes if s.shape_type == 13]
        self.assertEqual(len(pics), 11)
        texts = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        self.assertIn("BIG IDEAS - Q2 2027", texts)
        self.assertIn("PANTALONS", texts)
        self.assertEqual(pics[0].click_action.hyperlink.address, "https://x.example/p/0")
        self.assertIn("1. A", slide.notes_slide.notes_text_frame.text)

    def test_partial_and_unreadable_photos(self):
        from pptx import Presentation

        result, out = self.build(6, bad=("A0.jpg",))  # une photo illisible : 5 cases
        self.assertEqual((result["slides"], result["partial"]), (1, 1))
        slide = Presentation(str(out)).slides[0]
        self.assertEqual(len([s for s in slide.shapes if s.shape_type == 13]), 5)
        self.assertIn("incomplet", " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame))

    def test_crop_matches_cell_ratio(self):
        data = board.fetch_cell_image("https://x.example/a.jpg", (1.4, 2.1), getter=lambda u: jpeg(800, 400))
        w, h = Image.open(io.BytesIO(data)).size
        self.assertAlmostEqual(w / h, 1.4 / 2.1, places=1)
        self.assertIsNone(board.fetch_cell_image("https://x.example/a.jpg", (1, 1), getter=lambda u: b"pas une image"))


class VocabularyTests(unittest.TestCase):
    def test_normalize_new_attributes(self):
        raw = {"type": "Pantalons", "shape": "Barrel", "rise": "haute", "style": "Décontracté",
               "closure": "pull-on", "colors": ["blanc"]}
        out = analyze.normalize(raw, {"name": "Pantalon"}, True)
        self.assertEqual((out["shape"], out["rise"], out["style"], out["closure"]), ("barrel", "haute", "décontracté", "pull-on"))
        none = analyze.normalize(raw, {"name": "Pantalon"}, False)
        self.assertEqual((none["shape"], none["rise"], none["style"]), ("", "", ""))  # rien sans photo
        top = analyze.normalize({"type": "Hauts", "shape": "dolman", "rise": "haute", "closure": "pull-on"}, {"name": "T"}, True)
        self.assertEqual((top["shape"], top["rise"], top["closure"]), ("dolman", "", ""))
        wrong = analyze.normalize({"type": "Hauts", "shape": "barrel"}, {"name": "T"}, True)
        self.assertEqual(wrong["shape"], "")  # forme d'un autre type : écartée


class NewVocabularyTests(unittest.TestCase):
    def test_types_and_lengths_depend_on_type(self):
        short = analyze.normalize({"type": "Shorts", "length": "bermuda", "shape": "crayon"}, {"name": "x"}, True)
        self.assertEqual((short["type"], short["length"], short["shape"]), ("Shorts", "bermuda", ""))  # « crayon » n'est pas une forme de short
        pant = analyze.normalize({"type": "Pantalons", "length": "7/8", "shape": "palazzo"}, {"name": "x"}, True)
        self.assertEqual((pant["length"], pant["shape"]), ("7/8", "palazzo"))
        skirt = analyze.normalize({"type": "Jupes", "length": "bermuda", "shape": "crayon", "rise": "haute"}, {"name": "x"}, True)
        self.assertEqual((skirt["length"], skirt["shape"], skirt["rise"]), ("", "crayon", "haute"))  # bermuda n'existe pas pour une jupe
        jeans = analyze.normalize({"type": "Jeans", "shape": "baggy", "length": "traînant", "closure": "zippé ou boutonné"}, {"name": "x"}, True)
        self.assertEqual((jeans["shape"], jeans["length"], jeans["closure"]), ("baggy", "traînant", "zippé ou boutonné"))
        blazer = analyze.normalize({"type": "Blazers", "shape": "croisé", "rise": "haute"}, {"name": "x"}, True)
        self.assertEqual((blazer["shape"], blazer["rise"]), ("croisé", ""))

    def test_fallback_rules_split_the_big_categories(self):
        rules = {"Denim Short": "Shorts", "Linen Bermuda": "Shorts", "Short Sleeve Top": "Hauts", "Straight Leg Jeans": "Jeans",
                 "Tailored Blazer": "Blazers", "Wool Coat": "Manteaux", "Wide Leg Pant": "Pantalons", "Denim Jacket": "Manteaux"}
        for name, expected in rules.items():
            self.assertEqual(analyze.heuristic_type({"name": name}), expected, name)

    def test_details_have_english_twins(self):
        out = analyze.normalize({"type": "Robes", "details": ["Volants", "fente"], "details_en": ["Ruffles"]}, {"name": "x"}, True)
        self.assertEqual((out["details"], out["details_en"]), (["volants", "fente"], ["ruffles", ""]))

    def test_labels_cover_the_whole_vocabulary(self):
        import json

        data = json.loads((Path(__file__).resolve().parent.parent / "site" / "labels.json").read_text(encoding="utf-8"))["values"]
        words = set(analyze.TYPES) | set(analyze.COLORS) | set(analyze.MATERIALS) | set(analyze.CUTS) | set(analyze.NECKLINES)
        words |= set(analyze.RISES) | set(analyze.STYLES) | set(analyze.CLOSURES) | set(analyze.FIBRE_WORDS.values()) | {"mélange"}
        for vals in list(analyze.LENGTHS_BY_TYPE.values()) + list(analyze.SHAPES.values()):
            words |= set(vals)
        words |= {kind for kind, _ in analyze.TREND_KINDS}
        self.assertEqual(sorted(w for w in words if w not in data), [])

    def test_typed_trends(self):
        items = [{"id": f"{b}{n}", "brand": b, "type": "Jupes", "length": "midi", "shape": "crayon", "attrs_v": analyze.ATTRS_VERSION}
                 for b in "ABC" for n in range(2)]
        rows = analyze.compute_trends(items)
        self.assertIn("Jupes · midi", {r["value"] for r in rows if r["kind"] == "longueur"})
        self.assertIn("Jupes · crayon", {r["value"] for r in rows if r["kind"] == "forme"})

    def test_english_board_labels(self):
        board.LANG = "en"
        try:
            self.assertEqual(board.group_label("Jupes", ("crayon", "midi", "haute", "")), "SKIRTS · PENCIL · MIDI · HIGH-RISE")
            self.assertEqual(board.tu("Groupe incomplet : {n}/{m} produits. ", n=5, m=11), "Incomplete group: 5/11 products. ")
        finally:
            board.LANG = "fr"
        self.assertEqual(board.group_label("Jupes", ("crayon", "midi", "haute", "")), "JUPES · CRAYON · LONGUEUR MIDI · TAILLE HAUTE")

    def test_digest_has_english_version(self):
        class Client:
            class messages:
                @staticmethod
                def create(**kw):
                    class R:
                        stop_reason = "end_turn"
                        content = [type("B", (), {"type": "text", "text": '{"title":"Noir","summary":"Texte.","tags":["a"],"title_en":"Black","summary_en":"Text.","tags_en":["b"]}'})()]
                    return R
        stats = {"types": [], "palette": [], "materials": [], "lengths": [], "necklines": [], "details": [], "prices": {}, "trends": []}
        items = [{"brand": "A", "name": "x", "type": "Robes", "colors": [], "details": []}]
        digest = analyze.build_digest(Client, items, stats, items)
        self.assertEqual((digest["title"], digest["en"]["title"], digest["en"]["tags"]), ("Noir", "Black", ["b"]))
        self.assertEqual(analyze.week_label_en(__import__("datetime").date(2026, 10, 30)), "Week of October 26")


if __name__ == "__main__":
    unittest.main()
